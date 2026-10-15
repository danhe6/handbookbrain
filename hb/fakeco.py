"""Generate a complete FICTIONAL company ("Larkspur Freight") to test that HandbookBrain works on data it has
never seen, with no code changes - only companies/larkspur.yaml.

Deliberately messy and heterogeneous:
  docs     md (frontmatter) / Confluence-style html / pdf / docx / txt / csv, in dept folders, exported as
           biweekly dated snapshots (no git). Owners declared 5 different ways, or not at all.
  changes  material policy changes, cosmetic rewrites, heading renames, a removed doc, a new doc
  convos   Zoom .vtt + Zoom captions .txt, Teams .vtt, Slack export, .eml files, an .mbox mailing list,
           Jira CSV, meeting notes (human and AI-written; one AI note contains a planted wrong number)
  eval     questions.jsonl (gold answer + gold doc/section + type) and labels.json (event ground truth).
           The agent never reads data/companies/<co>/eval.

All people, companies and conversations are invented. Real-world facts are not asserted anywhere.
"""
import concurrent.futures as cf
import csv
import json
import random
import re
import shutil
import textwrap
from copy import deepcopy
from datetime import datetime, timedelta, timezone
from email.message import EmailMessage
from email.utils import format_datetime, make_msgid
from pathlib import Path

from .config import ROOT
from .llm import complete, usage_line

GEN_MODEL = __import__("os").getenv("HB_SYNTH_MODEL", "openai/gpt-4.1")
rnd = random.Random(7)

SYSTEM = """You create realistic internal documents and workplace communication for a FICTIONAL company used as a
test dataset. Write like real employees: specific numbers, names of (fictional) vendors, mild inconsistency in style,
occasional typos in chat. Never mention that anything is fictional. Return JSON only."""


def _llm(prompt: str, max_tokens: int = 5000) -> dict:
    r = complete(GEN_MODEL, [{"role": "system", "content": SYSTEM}, {"role": "user", "content": prompt}],
                 response_format={"type": "json_object"}, max_tokens=max_tokens)
    return json.loads(r.choices[0].message.content)


def _pmap(fn, items, workers=6):
    with cf.ThreadPoolExecutor(workers) as ex:
        return list(ex.map(fn, items))


SNAPSHOTS = [datetime(2026, 4, 15) + timedelta(days=14 * i) for i in range(13)]  # 2026-04-15 .. 2026-09-30
FORMATS = ["md"] * 6 + ["html"] * 5 + ["pdf"] * 4 + ["docx"] * 4 + ["txt"] * 2
OWNER_STYLES = {"md": ["frontmatter", "inline", "none"], "html": ["meta", "inline", "none"],
                "pdf": ["inline", "none", "author"], "docx": ["author", "inline", "none"], "txt": ["inline", "none"]}


# ================================================================== 1. company bible + documents

def plan_company() -> dict:
    return _llm("""Invent a mid-size company: "Larkspur Freight", a 160-person freight logistics company (HQ Denver,
offices in Rotterdam and Manila, plus remote staff). Produce:
- 20 employees: name, email (first.last@larkspur.example), title, team. Teams: People Ops, Finance, IT & Security,
  Operations, Legal, Engineering, Sales. Include 3-4 people per team for the first five teams.
- 22 internal policy/process documents across folders hr/, finance/, it-security/, operations/, legal/, engineering/.
  For each: id (kebab), title, folder, owner_name (one of the employees, or null for 4 of them), summary (1 line).
  Topics must be concrete and fact-heavy: PTO, parental leave, remote work stipend, travel & expenses, per diem,
  corporate card, procurement approvals, laptop & equipment, password/MFA, incident response, data retention,
  on-call compensation, driver safety, warehouse safety, vendor onboarding, code review, release process,
  learning budget, referral bonus, relocation, holiday calendar, offboarding.
Return JSON: {"people": [...], "docs": [...]}""", max_tokens=6000)


def write_doc_content(d: dict, company: dict) -> dict:
    out = _llm(f"""Write the content of this internal document for {company['name']}.
Title: {d['title']} (folder {d['folder']}). Summary: {d['summary']}.
5-8 sections. Each section: a heading and a body of 60-180 words with SPECIFIC facts (amounts with currency, day
counts, percentages, deadlines, approver roles, named fictional vendors/tools, eligibility rules). Bodies may contain
bullet lines starting with "- ". Different offices may have different rules.
Return JSON: {{"sections": [{{"heading": str, "body": str}}]}}""")
    d = dict(d)
    d["sections"] = [{"heading": s["heading"].strip(), "body": s["body"].strip()} for s in out["sections"]]
    return d


def write_csv_doc(company: dict) -> dict:
    out = _llm("""Create a per diem rate table for Larkspur Freight business travel: 14 cities across US, Netherlands,
Philippines, Germany, Mexico, Singapore. Columns: City, Country, Lodging cap (USD), Meals per day (USD), Notes.
Return JSON: {"header": [...], "rows": [[...], ...]}""")
    return {"id": "per-diem-rates", "title": "Per Diem Rates", "folder": "finance", "format": "csv",
            "owner_name": None, "owner_style": "none", "table": out}


# ================================================================== 2. change timeline (ground truth)

def plan_changes(docs: list[dict], people: list[dict]) -> list[dict]:
    """Code picks WHAT changes and WHEN; the LLM only writes the new text. Returns ground-truth events."""
    textual = [d for d in docs if d["format"] != "csv"]
    rnd.shuffle(textual)
    events = []
    # 8 material changes, decided somewhere, implemented later
    sources = ["zoom", "teams", "zoom_captions", "email", "mbox", "slack", "notes_ai", "none"]
    for i, d in enumerate(textual[:8]):
        sec = rnd.choice(d["sections"][1:] or d["sections"])
        impl = SNAPSHOTS[2] + timedelta(days=rnd.randint(0, 130))
        events.append({"kind": "material", "doc": d["id"], "heading": sec["heading"], "implemented_on": impl,
                       "decided_on": impl - timedelta(days=rnd.randint(5, 20)), "decided_in": sources[i]})
    for d in textual[8:10]:   # cosmetic rewrites: NOT a policy change
        sec = rnd.choice(d["sections"])
        events.append({"kind": "cosmetic", "doc": d["id"], "heading": sec["heading"],
                       "implemented_on": SNAPSHOTS[3] + timedelta(days=rnd.randint(0, 100))})
    d = textual[10]           # heading renamed, body unchanged
    sec = rnd.choice(d["sections"][1:])
    events.append({"kind": "rename", "doc": d["id"], "heading": sec["heading"],
                   "implemented_on": SNAPSHOTS[4] + timedelta(days=3)})
    events.append({"kind": "remove_doc", "doc": textual[11]["id"], "implemented_on": SNAPSHOTS[8] + timedelta(days=2)})
    events.append({"kind": "add_doc", "doc": "ai-tools-acceptable-use", "implemented_on": SNAPSHOTS[6] + timedelta(days=4),
                   "decided_on": SNAPSHOTS[6] - timedelta(days=6), "decided_in": "teams"})
    for d, src in zip(textual[12:15], ["zoom", "email", "teams"]):   # decided, never implemented
        sec = rnd.choice(d["sections"][1:])
        events.append({"kind": "unimplemented", "doc": d["id"], "heading": sec["heading"], "decided_in": src,
                       "decided_on": SNAPSHOTS[3] + timedelta(days=rnd.randint(0, 80))})

    by_id = {d["id"]: d for d in docs}

    def write(ev):
        if ev["kind"] in ("material", "unimplemented", "cosmetic"):
            d = by_id[ev["doc"]]
            body = next(s["body"] for s in d["sections"] if s["heading"] == ev["heading"])
            instr = ("Change 1-2 concrete facts (an amount, day count, percentage, approver, vendor or deadline) in a "
                     "realistic way; keep everything else." if ev["kind"] != "cosmetic" else
                     "Rewrite for clarity and formatting ONLY. Every fact, number, name and rule must stay exactly the same.")
            out = _llm(f"""Policy section "{ev['heading']}" in "{d['title']}":
{body}

{instr}
Return JSON: {{"new_body": str, "change_summary": "what changed, old -> new (or 'no factual change')",
 "question": "a natural employee question whose answer depends on the changed fact",
 "old_answer": "short answer per OLD text", "new_answer": "short answer per NEW text"}}""")
            ev.update(out)
        elif ev["kind"] == "rename":
            out = _llm(f"""Give a different but equivalent heading for the section "{ev['heading']}" (same meaning).
Return JSON: {{"new_heading": str}}""")
            ev.update(out)
        elif ev["kind"] == "add_doc":
            d = write_doc_content({"id": ev["doc"], "title": "AI Tools Acceptable Use Policy", "folder": "it-security",
                                   "summary": "Which AI assistants employees may use, with what data, and approval steps"},
                                  {"name": "Larkspur Freight"})
            ev["new_doc"] = d
        return ev

    events = _pmap(write, events)
    for e in events:
        for k in ("implemented_on", "decided_on"):
            if isinstance(e.get(k), datetime):
                e[k] = e[k].replace(hour=15).isoformat()
        e["decided_by"] = rnd.choice(people)["name"]
    return events


def doc_state(doc: dict, events: list[dict], at: datetime) -> dict | None:
    """The document as it exists at time `at`, after applying every event implemented by then."""
    d = deepcopy(doc)
    for e in events:
        if e["doc"] != d["id"] or not e.get("implemented_on") or datetime.fromisoformat(e["implemented_on"]) > at:
            continue
        if e["kind"] == "remove_doc":
            return None
        for s in d.get("sections", []):
            if s["heading"] == e.get("heading"):
                if e["kind"] in ("material", "cosmetic"):
                    s["body"] = e["new_body"]
                elif e["kind"] == "rename":
                    s["heading"] = e["new_heading"]
    return d


# ================================================================== 3. renderers (one per format)

def _paras(body: str) -> list[tuple[str, str]]:
    out = []
    for line in body.split("\n"):
        line = line.strip()
        if line:
            out.append(("li", line[2:]) if line.startswith("- ") else ("p", line))
    return out


def _esc(t: str) -> str:
    return str(t).replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")


def _owner_line(d):
    return f"Document owner: {d['owner_name']}" if d.get("owner_name") else ""


def render_md(d, people):
    fm = {"title": d["title"], "last_reviewed": "2026"}
    if d["owner_style"] == "frontmatter" and d.get("owner_name"):
        fm["owner"] = d["owner_name"]
    out = ["---"] + [f"{k}: {json.dumps(v)}" for k, v in fm.items()] + ["---", "", f"# {d['title']}", ""]
    if d["owner_style"] == "inline" and d.get("owner_name"):
        out += [f"**Owner:** {d['owner_name']}", ""]
    for s in d["sections"]:
        out += [f"## {s['heading']}", "", s["body"], ""]
    return "\n".join(out).encode()


def render_html(d, people):
    meta = f'<meta name="owner" content="{d["owner_name"]}">' if d["owner_style"] == "meta" and d.get("owner_name") else ""
    body = [f'<div id="breadcrumbs">Spaces / Larkspur Internal / {d["folder"].title()}</div>',
            '<nav>Home · Recent · Spaces · People</nav>', f"<h1>{_esc(d['title'])}</h1>"]
    if d["owner_style"] == "inline" and d.get("owner_name"):
        body.append(f"<p><strong>Page owner:</strong> {_esc(d['owner_name'])}</p>")
    for s in d["sections"]:
        body.append(f"<h2>{_esc(s['heading'])}</h2>")
        items = _paras(s["body"])
        lis = [t for k, t in items if k == "li"]
        body += [f"<p>{_esc(t)}</p>" for k, t in items if k == "p"]
        if lis:
            body.append("<ul>" + "".join(f"<li>{_esc(t)}</li>" for t in lis) + "</ul>")
    body.append('<footer>Last updated by Confluence export · Like · Watch · Share</footer>')
    return (f"<html><head><title>{d['title']} - Larkspur Internal</title>{meta}</head><body>"
            + "\n".join(body) + "</body></html>").encode()


def render_pdf(d, people):
    import io
    from reportlab.lib.pagesizes import letter
    from reportlab.lib.styles import getSampleStyleSheet
    from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer
    buf = io.BytesIO()
    author = d["owner_name"] if d["owner_style"] == "author" and d.get("owner_name") else ""

    def header(c, doc):
        c.saveState()
        c.setFont("Helvetica", 8)
        c.drawString(72, 760, "Larkspur Freight - Internal - Do not distribute")
        c.drawString(520, 30, str(doc.page))
        c.restoreState()

    doc = SimpleDocTemplate(buf, pagesize=letter, title=d["title"], author=author)
    st = getSampleStyleSheet()
    flow = [Paragraph(_esc(d["title"]), st["Title"])]
    if d["owner_style"] == "inline" and d.get("owner_name"):
        flow.append(Paragraph(_esc(_owner_line(d)), st["Normal"]))
    for i, s in enumerate(d["sections"], 1):
        flow += [Spacer(1, 8), Paragraph(f"{i}. {_esc(s['heading'])}", st["Heading2"])]
        for k, t in _paras(s["body"]):
            flow.append(Paragraph(("&bull; " if k == "li" else "") + _esc(t), st["Normal"]))
    doc.build(flow, onFirstPage=header, onLaterPages=header)
    return buf.getvalue()


def render_docx(d, people):
    import io
    import docx
    doc = docx.Document()
    doc.core_properties.title = d["title"]
    doc.core_properties.author = d["owner_name"] if d["owner_style"] == "author" and d.get("owner_name") else ""
    doc.add_heading(d["title"], level=0)
    if d["owner_style"] == "inline" and d.get("owner_name"):
        doc.add_paragraph(_owner_line(d))
    for s in d["sections"]:
        doc.add_heading(s["heading"], level=1)
        for k, t in _paras(s["body"]):
            doc.add_paragraph(t, style="List Bullet" if k == "li" else None)
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def render_txt(d, people):
    out = [d["title"], "=" * len(d["title"]), ""]
    if d["owner_style"] == "inline" and d.get("owner_name"):
        out += [_owner_line(d), ""]
    for s in d["sections"]:
        out += [s["heading"].upper(), ""]
        for k, t in _paras(s["body"]):
            out += textwrap.wrap(("* " if k == "li" else "") + t, 92) + [""]
    return "\n".join(out).encode()


def render_csv(d, people):
    import io
    buf = io.StringIO()
    w = csv.writer(buf)
    w.writerow(d["table"]["header"])
    w.writerows(d["table"]["rows"])
    return buf.getvalue().encode()


RENDER = {"md": render_md, "html": render_html, "pdf": render_pdf, "docx": render_docx, "txt": render_txt,
          "csv": render_csv}
EXT = {"md": ".md", "html": ".html", "pdf": ".pdf", "docx": ".docx", "txt": ".txt", "csv": ".csv"}


def doc_path(d) -> str:
    return f"{d['folder']}/{d['id']}{EXT[d['format']]}"


def write_snapshots(out: Path, docs: list[dict], events: list[dict], people):
    added = {e["doc"]: e for e in events if e["kind"] == "add_doc"}
    for snap in SNAPSHOTS:
        sd = out / snap.strftime("%Y-%m-%d")
        sd.mkdir(parents=True, exist_ok=True)
        for d in docs:
            st = doc_state(d, events, snap)
            if st:
                (sd / doc_path(st)).parent.mkdir(parents=True, exist_ok=True)
                (sd / doc_path(st)).write_bytes(RENDER[st["format"]](st, people))
        for e in added.values():
            if datetime.fromisoformat(e["implemented_on"]) <= snap:
                nd = {**e["new_doc"], "format": "md", "owner_style": "frontmatter",
                      "owner_name": e["decided_by"]}
                (sd / doc_path(nd)).parent.mkdir(parents=True, exist_ok=True)
                (sd / doc_path(nd)).write_bytes(render_md(nd, people))


# ================================================================== 4. conversations

def _vtt_t(sec):
    h, r = divmod(sec, 3600)
    m, s = divmod(r, 60)
    return f"{int(h):02d}:{int(m):02d}:{s:06.3f}"


def _talk(prompt, n=(30, 45)) -> list[dict]:
    out = _llm(f"""{prompt}
Write {rnd.randint(*n)} turns. Return JSON: {{"turns": [{{"speaker": name, "text": str}}]}}""", max_tokens=6000)
    return [t for t in out.get("turns", []) if t.get("text")]


def write_meeting(root: Path, kind: str, mid: str, title: str, when: datetime, attendees: list[str], turns: list[dict]):
    d = root / ("zoom" if kind in ("zoom", "zoom_captions") else "teams") / mid
    d.mkdir(parents=True, exist_ok=True)
    t, cues = 3.0, []
    for u in turns:
        dur = max(1.5, len(u["text"].split()) / 2.7)
        cues.append((t, t + dur, u))
        t += dur + rnd.uniform(0.3, 1.8)
    if kind == "zoom":
        lines = ["WEBVTT", ""]
        for i, (a, b, u) in enumerate(cues, 1):
            lines += [str(i), f"{_vtt_t(a)} --> {_vtt_t(b)}", f"{u['speaker']}: {u['text']}", ""]
        (d / "transcript.vtt").write_text("\n".join(lines))
    elif kind == "zoom_captions":   # in-meeting "Save captions" export
        lines = []
        for a, b, u in cues:
            ts = when + timedelta(seconds=a)
            lines += [f"[{u['speaker']}] {ts:%H:%M:%S}", u["text"]]
        (d / "captions.txt").write_text("\n".join(lines))
    else:
        lines = ["WEBVTT", ""]
        for a, b, u in cues:
            lines += [f"{_vtt_t(a)} --> {_vtt_t(b)}", f"<v {u['speaker']}>{u['text']}</v>", ""]
        (d / "transcript.vtt").write_text("\n".join(lines))
    start = when.replace(tzinfo=timezone.utc).strftime("%Y-%m-%dT%H:%M:%SZ")
    meta = ({"uuid": mid, "topic": title, "start_time": start, "participants": [{"name": n} for n in attendees]}
            if kind != "teams" else
            {"id": mid, "subject": title, "startDateTime": start,
             "participants": {"attendees": [{"identity": {"user": {"displayName": n}}} for n in attendees]}})
    (d / "meeting.json").write_text(json.dumps({**meta, "synthetic": True}, indent=2))
    return f"{'teams' if kind == 'teams' else 'zoom'}:{mid}"


def _bt(day: datetime) -> datetime:
    while day.weekday() >= 5:
        day -= timedelta(days=1)
    return day.replace(hour=rnd.choice([14, 15, 16, 17, 19]), minute=rnd.choice([0, 30]), second=0, microsecond=0)


class Slack:
    def __init__(self, people):
        self.ids = {p["name"]: f"U{i:04d}" for i, p in enumerate(people)}
        self.people = people
        self.days: dict = {}

    def thread(self, ch, start, msgs) -> str:
        t, root = start, None
        for i, m in enumerate(msgs):
            uid = self.ids.get(m["speaker"]) or rnd.choice(list(self.ids.values()))
            ts = f"{t.replace(tzinfo=timezone.utc).timestamp():.6f}"
            rec = {"type": "message", "user": uid, "text": m["text"], "ts": ts, "synthetic": True}
            if i == 0:
                root = ts
                rec.update(thread_ts=ts, reply_count=len(msgs) - 1)
            else:
                rec.update(thread_ts=root)
            self.days.setdefault((ch, t.strftime("%Y-%m-%d")), []).append(rec)
            t += timedelta(minutes=rnd.randint(3, 120))
        return f"slack:{ch}:{root}"

    def flush(self, root: Path):
        d = root / "slack"
        d.mkdir(parents=True, exist_ok=True)
        (d / "users.json").write_text(json.dumps([{"id": i, "name": n.split()[0].lower(), "real_name": n}
                                                  for n, i in self.ids.items()], indent=2))
        (d / "channels.json").write_text(json.dumps([{"name": c} for c in sorted({c for c, _ in self.days})]))
        for (ch, day), ms in self.days.items():
            (d / ch).mkdir(exist_ok=True)
            (d / ch / f"{day}.json").write_text(json.dumps(sorted(ms, key=lambda m: float(m["ts"])), indent=2))


def write_email_thread(root: Path, msgs: list[dict], subject: str, start: datetime, people, mbox: bool) -> str:
    from .adapters import _slug, utc_iso
    email_of = {p["name"]: p["email"] for p in people}
    out, t, prev = [], start, None
    for i, m in enumerate(msgs):
        e = EmailMessage()
        e["From"] = f'{m["speaker"]} <{email_of.get(m["speaker"], "someone@larkspur.example")}>'
        e["To"] = m.get("to") or "all-staff@larkspur.example"
        e["Subject"] = ("Re: " if i else "") + subject
        e["Date"] = format_datetime(t.replace(tzinfo=timezone.utc))
        e["Message-ID"] = make_msgid(domain="larkspur.example")
        if prev:
            e["In-Reply-To"] = prev
        e["X-Synthetic"] = "true"
        quoted = "\n\nOn an earlier date someone wrote:\n> (previous message)" if i else ""
        e.set_content(m["text"] + quoted)
        prev = e["Message-ID"]
        out.append(e)
        t += timedelta(hours=rnd.randint(1, 30))
    d = root / "email"
    d.mkdir(parents=True, exist_ok=True)
    if mbox:
        with open(d / "people-ops-announce.mbox", "a") as f:
            for e in out:
                f.write(f"From MAILER-DAEMON {start:%a %b %d %H:%M:%S %Y}\n")
                f.write(e.as_string().replace("\nFrom ", "\n>From ") + "\n\n")
    else:
        for i, e in enumerate(out):
            (d / f"{_slug(subject)[:40]}-{start:%Y%m%d}-{i}.eml").write_bytes(bytes(e))
    return f"email:{_slug(subject)}:{utc_iso(start.replace(tzinfo=timezone.utc))}"


def write_notes(root: Path, name: str, title: str, when: datetime, body: str, ai: bool, attendees, author=None):
    d = root / "notes"
    d.mkdir(parents=True, exist_ok=True)
    fm = {"title": title, "date": when.replace(tzinfo=timezone.utc).isoformat(), "attendees": attendees,
          "synthetic": True, **({"generated_by": "ai"} if ai else {"author": author})}
    (d / f"{name}.md").write_text("---\n" + "\n".join(f"{k}: {json.dumps(v)}" for k, v in fm.items()) + "\n---\n\n" + body)
    return f"notes:{name}"


def gen_conversations(root: Path, docs, events, people) -> list[dict]:
    by_id = {d["id"]: d for d in docs}
    names = [p["name"] for p in people]
    team_of = {p["name"]: p["team"] for p in people}
    slack = Slack(people)
    labels = []

    def roster(team, k=4):
        same = [n for n in names if team_of[n] == team]
        return (same + rnd.sample([n for n in names if n not in same], k))[:k]

    def decision_conv(ev):
        d = by_id.get(ev["doc"]) or ev.get("new_doc")
        when = _bt(datetime.fromisoformat(ev["decided_on"]))
        owner_team = next((p["team"] for p in people if p["name"] == d.get("owner_name")), "People Ops")
        att = list(dict.fromkeys([ev["decided_by"]] + roster(owner_team)))
        if ev["kind"] == "add_doc":
            what = "adopt a NEW 'AI Tools Acceptable Use Policy' (which assistants are allowed, with which data, approvals)"
        else:
            what = (f"change the policy '{d['title']} > {ev['heading']}': {ev['change_summary']}. "
                    f"Current text:\n{next(s['body'] for s in d['sections'] if s['heading'] == ev['heading'])[:1200]}")
        src = ev["decided_in"]
        lab = {"scenario": "decision_implemented" if ev["kind"] != "unimplemented" else "decision_unimplemented",
               "doc": doc_path(d) if ev["kind"] != "add_doc" else f"it-security/{ev['doc']}.md",
               "heading": ev.get("heading"), "decided_on": ev["decided_on"], "implemented_on": ev.get("implemented_on"),
               "decision": ev.get("change_summary") or what, "decided_by": ev["decided_by"], "source": src}
        if src in ("zoom", "teams", "zoom_captions"):
            turns = _talk(f"""Meeting "{owner_team} weekly" at Larkspur Freight on {when:%A %B %d, %Y}. Attendees: {att}.
Start with small talk and two quick status updates. Then the group discusses and DECIDES to {what}.
Someone pushes back once; {ev['decided_by']} states the final decision clearly; someone takes the action item to
update the document. End with one unrelated item (e.g. a vendor renewal or hiring update).""")
            mid = f"{src}-{when:%Y%m%d}-{ev['doc'][:20]}"
            lab["cid"] = write_meeting(root, src, mid, f"{owner_team} weekly", when, att, turns)
        elif src in ("email", "mbox"):
            msgs = _llm(f"""Email thread at Larkspur Freight starting {when:%B %d, %Y}. {ev['decided_by']} announces the decision
to {what}, with effective date and rationale. 2-3 replies from {att[1:]} asking clarifying questions or
acknowledging. Return JSON: {{"subject": str, "messages": [{{"speaker": name, "text": str}}]}}""")
            lab["cid"] = write_email_thread(root, msgs["messages"], msgs["subject"], when, people, mbox=src == "mbox")
        elif src == "slack":
            msgs = _talk(f"""Slack thread in #{owner_team.lower().replace(' ', '-').replace('&', 'and')}-leads. {ev['decided_by']} and 2 colleagues
quickly agree to {what}. {ev['decided_by']} says "ok decided, I'll update the doc".""", (5, 9))
            lab["cid"] = slack.thread(owner_team.lower().replace(" & ", "-").replace(" ", "-") + "-leads", when, msgs)
        elif src == "notes_ai":
            body = _llm(f"""AI meeting-notetaker summary of "{owner_team} sync" on {when:%B %d, %Y} with {att}.
Sections: Summary, Decisions, Action items, Other topics. Decision: {what}.
IMPORTANT: in the Decisions section, state the NEW value with ONE number slightly wrong (e.g. 15 instead of 10),
as AI notetakers sometimes mishear. Return JSON: {{"markdown": str, "planted_error": "what is wrong"}}""")
            lab["planted_error"] = body["planted_error"]
            lab["cid"] = write_notes(root, f"ai-notes-{when:%Y%m%d}-{ev['doc'][:20]}", f"{owner_team} sync (AI notes)",
                                     when, body["markdown"], True, att)
        else:
            lab["cid"] = None   # change made with no recorded decision
        return lab

    dec_events = [e for e in events if e["kind"] in ("material", "unimplemented", "add_doc")]
    labels += _pmap(decision_conv, dec_events)

    # conflicts: someone states the OLD rule in Slack after the change shipped
    def conflict(ev):
        d = by_id[ev["doc"]]
        when = _bt(datetime.fromisoformat(ev["implemented_on"]) + timedelta(days=rnd.randint(4, 25)))
        a, b = rnd.sample([n for n in names if team_of[n] in ("Operations", "Engineering", "Sales")] or names, 2)
        msgs = _talk(f"""Slack thread in #ask-people-ops. {a} asks: "{ev['question']}". {b} answers confidently with the OLD
rule: "{ev['old_answer']}". {a} thanks them. Nobody corrects it.""", (3, 5))
        return {"scenario": "conflict", "cid": slack.thread("ask-people-ops", when, msgs), "doc": doc_path(d),
                "heading": ev["heading"], "old": ev["old_answer"], "new": ev["new_answer"]}

    labels += _pmap(conflict, [e for e in events if e["kind"] == "material"][:6])

    # answerable / tribal / noise threads, gap tickets in Jira
    stable = [(d, s) for d in docs if d["format"] != "csv" for s in d["sections"]
              if not any(e.get("doc") == d["id"] and e.get("heading") == s["heading"] for e in events)]
    rnd.shuffle(stable)
    gap_topics = ["whether pets are allowed in the Rotterdam office", "reimbursing a home internet upgrade above the stipend",
                  "taking a sabbatical after 5 years", "using personal phones for warehouse scanning apps",
                  "getting a parking pass at the Denver HQ"]

    def answerable(ds):
        d, s = ds
        when = _bt(SNAPSHOTS[0] + timedelta(days=rnd.randint(5, 160)))
        a, b = rnd.sample(names, 2)
        msgs = _talk(f"""Slack thread in #ask-people-ops. {a} asks a practical question answered by this policy
"{d['title']} > {s['heading']}": {s['body'][:700]}
{b} replies roughly correctly and says it's in the {d['title']} doc.""", (3, 5))
        return {"scenario": "answerable", "cid": slack.thread("ask-people-ops", when, msgs), "doc": doc_path(d),
                "heading": s["heading"]}

    def tribal(topic):
        when = _bt(SNAPSHOTS[0] + timedelta(days=rnd.randint(5, 160)))
        a, b = rnd.choice(names), rnd.choice([n for n in names if team_of[n] == "People Ops"] or names)
        msgs = _talk(f"""Slack thread in #ask-people-ops. {a} asks about {topic}. {b} answers with a specific unwritten practice
(an amount, approver or form) and adds "this isn't written down anywhere yet".""", (3, 5))
        return {"scenario": "tribal", "cid": slack.thread("ask-people-ops", when, msgs), "topic": topic}

    def noise(_):
        when = _bt(SNAPSHOTS[0] + timedelta(days=rnd.randint(5, 160)))
        msgs = _talk("Slack thread in #random: casual chat (lunch spots in Denver, a sports game, a funny warehouse cat).", (3, 6))
        return {"scenario": "noise", "cid": slack.thread("random", when, msgs)}

    labels += _pmap(answerable, stable[:8]) + _pmap(tribal, gap_topics[:3]) + _pmap(noise, range(5))

    # Jira helpdesk export: gap topics asked repeatedly + answerable
    def ticket(i_spec):
        i, (kind, x) = i_spec
        about = (f"{x} (not covered by any policy)" if kind == "gap" else
                 f"something answered by '{x[0]['title']} > {x[1]['heading']}': {x[1]['body'][:500]}")
        t = _llm(f"""A Larkspur Freight employee files an internal helpdesk ticket (Jira, project PEOPLE) about {about}.
Phrase it in your own words. 0-2 comments from People Ops. Return JSON: {{"summary": str, "description": str,
"reporter": name, "comments": [{{"author": name, "body": str}}]}}""")
        when = _bt(SNAPSHOTS[0] + timedelta(days=rnd.randint(3, 165)))
        return i, kind, x, t, when

    specs = [("gap", g) for g in gap_topics for _ in range(rnd.randint(2, 3))] + [("answerable", ds) for ds in stable[8:16]]
    rnd.shuffle(specs)
    rows = []
    for i, kind, x, t, when in _pmap(ticket, list(enumerate(specs))):
        key = f"PEOPLE-{100 + i}"
        cm = [f"{(when + timedelta(hours=4 + 3 * j)):%d/%b/%y %I:%M %p};{c['author']};{c['body']}"
              for j, c in enumerate(t.get("comments", [])[:2])]
        rows.append([key, t["summary"], t["description"], t["reporter"], f"{when:%d/%b/%y %I:%M %p}",
                     rnd.choice(["Open", "In Progress", "Done"]), "synthetic"] + cm + [""] * (2 - len(cm)))
        labels.append({"scenario": kind, "cid": f"jira:{key}", **({"topic": x} if kind == "gap" else
                       {"doc": doc_path(x[0]), "heading": x[1]["heading"]})})
    (root / "jira").mkdir(parents=True, exist_ok=True)
    with open(root / "jira" / "PEOPLE-export.csv", "w", newline="") as f:
        w = csv.writer(f)
        w.writerow(["Issue key", "Summary", "Description", "Reporter", "Created", "Status", "Labels", "Comment", "Comment"])
        w.writerows(rows)

    # human-written notes with a decision that is just a to-do (should not become a policy alert)
    when = _bt(SNAPSHOTS[5])
    att = rnd.sample(names, 5)
    body = _llm(f"""Human-written meeting notes (markdown) from the "Q3 offsite planning" meeting on {when:%B %d, %Y}, attendees {att}.
Include: venue options, budget, a decision to book the venue in Boulder, action items. No policy changes.
Return JSON: {{"markdown": str}}""")["markdown"]
    labels.append({"scenario": "noise", "cid": write_notes(root, f"offsite-planning-{when:%Y%m%d}", "Q3 offsite planning",
                                                           when, body, False, att, author=att[0])})
    slack.flush(root)
    return labels


# ================================================================== 5. evaluation questions (gold)

def gen_questions(docs, events, people) -> list[dict]:
    by_id = {d["id"]: d for d in docs}
    qs = []
    changed = {(e["doc"], e.get("heading")) for e in events}
    stable = [(d, s) for d in docs if d["format"] != "csv" for s in d["sections"] if (d["id"], s["heading"]) not in changed
              and not any(e["kind"] == "remove_doc" and e["doc"] == d["id"] for e in events)]
    rnd.shuffle(stable)

    def lookup(ds):
        d, s = ds
        out = _llm(f"""From this policy section write one specific employee question and its short gold answer.
"{d['title']} > {s['heading']}": {s['body']}
Return JSON: {{"question": str, "answer": str}}""")
        return {"type": "lookup", "ask_at": "end", "q": out["question"], "gold": out["answer"],
                "doc": doc_path(d), "heading": s["heading"], "format": d["format"]}

    qs += _pmap(lookup, stable[:12])
    for e in events:
        d = by_id.get(e["doc"])
        if e["kind"] == "material":
            # asked at T0 (should go stale when the change ships), and again at the end (must give the NEW answer)
            qs.append({"type": "changed_fact", "ask_at": "t0", "q": e["question"], "gold": e["old_answer"],
                       "gold_after": e["new_answer"], "doc": doc_path(d), "heading": e["heading"],
                       "expect_stale_on": e["implemented_on"], "format": d["format"]})
            if e["decided_in"] != "none":
                qs.append({"type": "why_changed", "ask_at": "end",
                           "q": f"Why did the {e['heading'].lower()} rule in the {d['title']} change, and who decided it?",
                           "gold": f"{e['change_summary']}; decided by {e['decided_by']} via {e['decided_in']} around "
                                   f"{e['decided_on'][:10]}", "doc": doc_path(d), "heading": e["heading"]})
        elif e["kind"] == "cosmetic":
            qs.append({"type": "cosmetic_control", "ask_at": "t0", "q": e["question"], "gold": e["old_answer"],
                       "doc": doc_path(d), "heading": e["heading"], "expect_stale_on": None, "format": d["format"]})
    for d in [d for d in docs if d.get("owner_name") and d["owner_style"] != "none"][:4]:
        qs.append({"type": "who_owns", "ask_at": "end", "q": f"Who owns the {d['title']} document?",
                   "gold": d["owner_name"], "doc": doc_path(d), "heading": None})
    csvd = next(d for d in docs if d["format"] == "csv")
    row = rnd.choice(csvd["table"]["rows"])
    qs.append({"type": "table_lookup", "ask_at": "end",
               "q": f"What is the meals per diem in {row[0]}?", "gold": f"{row[3]} USD per day",
               "doc": doc_path(csvd), "heading": None, "format": "csv"})
    for q in ["Does Larkspur reimburse pet insurance?", "What is the company policy on bringing pets to the Rotterdam office?",
              "How many sabbatical weeks do employees get after 5 years?"]:
        qs.append({"type": "unanswerable", "ask_at": "end", "q": q, "gold": "Not covered by company documents.",
                   "doc": None, "heading": None})
    for i, q in enumerate(qs):
        q["id"] = f"q{i + 1:02d}"
    return qs


# ================================================================== main

YAML = """# Onboarding a company = this file. No code changes.
name: {name_title}
graph: {name}
sources:                       # several sources; higher authority wins when they disagree
  - name: wiki
    label: Confluence wiki (daily export)
    type: snapshots            # dated export folders <path>/YYYY-MM-DD/..., any format (md/html/pdf/docx/txt/csv)
    path: data/companies/{name}/docs
    kind: live                 # re-synced daily (hb sync)
    authority: 2
  - name: drive
    label: Shared drive
    type: uploads              # static files dropped into <path>/<upload date>/... (hb upload)
    path: data/companies/{name}/drive
    kind: static
    authority: 1
t0: "{t0}"
conversations: data/companies/{name}/conversations   # zoom/ teams/ slack/ email/ jira/ notes/
labels: data/companies/{name}/eval/labels.json        # ground truth, used only by `hb bench`
questions: data/companies/{name}/eval/questions.jsonl
"""


def generate(name: str = "larkspur") -> dict:
    base = ROOT / "data" / "companies" / name
    if base.exists():
        shutil.rmtree(base)
    docs_dir, conv_dir, eval_dir = base / "docs", base / "conversations", base / "eval"
    for p in (docs_dir, conv_dir, eval_dir):
        p.mkdir(parents=True)

    plan = plan_company()
    people = plan["people"]
    company = {"name": "Larkspur Freight"}
    specs = plan["docs"][:21]
    fmts = FORMATS[:]
    rnd.shuffle(fmts)
    for d, f in zip(specs, fmts):
        d["format"] = f
        d["owner_style"] = rnd.choice(OWNER_STYLES[f]) if d.get("owner_name") else "none"
        d["folder"] = d.get("folder", "hr").strip("/").split("/")[-1]
    docs = _pmap(lambda d: write_doc_content(d, company), specs) + [write_csv_doc(company)]
    print(f"docs: {len(docs)} {sorted({d['format'] for d in docs})}")

    events = plan_changes(docs, people)
    print(f"events: { {k: sum(e['kind'] == k for e in events) for k in {e['kind'] for e in events}} }")
    write_snapshots(docs_dir, docs, events, people)
    labels = gen_conversations(conv_dir, docs, events, people)
    questions = gen_questions(docs, events, people)

    (eval_dir / "company.json").write_text(json.dumps({"people": people, "docs": docs}, indent=2, default=str))
    (eval_dir / "events.json").write_text(json.dumps(events, indent=2, default=str))
    (eval_dir / "labels.json").write_text(json.dumps({"synthetic": True, "items": labels}, indent=2, default=str))
    with open(eval_dir / "questions.jsonl", "w") as f:
        for q in questions:
            f.write(json.dumps(q) + "\n")
    (ROOT / "companies").mkdir(exist_ok=True)
    (ROOT / "companies" / f"{name}.yaml").write_text(YAML.format(name=name, name_title=("Larkspur Freight" if name == "larkspur" else name.title()), t0=f"{SNAPSHOTS[0]:%Y-%m-%d}"))
    (base / "README.md").write_text("Everything under this folder is SYNTHETIC: a fictional company generated by hb/fakeco.py "
                                    "to test HandbookBrain on unseen data. eval/ is ground truth and is never read by the agent.\n")
    summary = {"docs": len(docs), "snapshots": len(SNAPSHOTS), "events": len(events), "conversation_labels": len(labels),
               "questions": len(questions)}
    print(summary, usage_line())
    return summary
