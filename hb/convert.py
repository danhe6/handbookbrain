"""Any document format -> normalized markdown (headings as #) + metadata (title, owner, owner_source).
Downstream (sectioning, versioning, graph) only ever sees markdown, so adding a format = one function here."""
import csv
import io
import re
from html.parser import HTMLParser
from pathlib import Path

import yaml

OWNER_LINE = re.compile(r"^\s*(?:\*\*)?(?:document\s+|policy\s+|page\s+)?(?:owner|maintainer|dri)(?:\*\*)?\s*[:\-–]\s*(.+?)\s*$",
                        re.I | re.M)


def _owner_from_text(text: str) -> str | None:
    m = OWNER_LINE.search(text[:3000])
    return m.group(1).strip(" *_`") if m else None


# ------------------------------------------------------------------ per format

def _md(raw: str) -> tuple[str, dict]:
    meta = {}
    if raw.startswith("---"):
        parts = raw.split("---", 2)
        if len(parts) == 3:
            try:
                meta = yaml.safe_load(parts[1]) or {}
            except yaml.YAMLError:
                meta = {}
    return raw, {"title": meta.get("title"), "owner": meta.get("owner") or meta.get("maintainer")}


def _txt(raw: str) -> tuple[str, dict]:
    """Plain text: ALL-CAPS lines, lines underlined with ===/---, and numbered '1. Title' lines become headings."""
    lines, out = raw.splitlines(), []
    title = None
    for i, l in enumerate(lines):
        s = l.strip()
        nxt = lines[i + 1].strip() if i + 1 < len(lines) else ""
        if re.fullmatch(r"[=\-]{3,}", s):
            continue
        if s and re.fullmatch(r"[=]{3,}", nxt):
            out.append(f"# {s}")
            title = title or s
        elif s and re.fullmatch(r"[-]{3,}", nxt) or (s.isupper() and 3 < len(s) < 70 and re.search(r"[A-Z]{3}", s)):
            out.append(f"## {s.title() if s.isupper() else s}")
        elif re.fullmatch(r"\d+(\.\d+)*\.?\s+[A-Z][^.]{2,70}", s):
            out.append(f"## {s}")
        else:
            out.append(l)
    return "\n".join(out), {"title": title}


class _HTML(HTMLParser):
    BLOCK = {"p", "div", "li", "tr", "br", "section", "article", "table"}

    def __init__(self):
        super().__init__()
        self.out, self.meta, self.cur, self.skip, self.in_title = [], {}, "", 0, False
        self.title = None

    def handle_starttag(self, tag, attrs):
        a = dict(attrs)
        if tag in ("script", "style", "nav", "footer"):
            self.skip += 1
        elif tag == "meta" and a.get("name", "").lower() in ("owner", "author", "maintainer"):
            self.meta[a["name"].lower()] = a.get("content")
        elif tag == "title":
            self.in_title = True
        elif re.fullmatch(r"h[1-6]", tag):
            self._flush()
            self.cur = "#" * min(int(tag[1]), 3) + " "
        elif tag == "li":
            self._flush()
            self.cur = "- "
        elif tag in ("td", "th"):
            self.cur += " | "

    def handle_endtag(self, tag):
        if tag in ("script", "style", "nav", "footer"):
            self.skip = max(0, self.skip - 1)
        elif tag == "title":
            self.in_title = False
        elif re.fullmatch(r"h[1-6]", tag) or tag in self.BLOCK:
            self._flush()

    def handle_data(self, data):
        if self.in_title:
            self.title = (self.title or "") + data.strip()
        elif not self.skip:
            self.cur += re.sub(r"\s+", " ", data)

    def _flush(self):
        if self.cur.strip(" |-#"):
            self.out.append(self.cur.strip())
        self.cur = ""


def _html(raw: str) -> tuple[str, dict]:
    p = _HTML()
    p.feed(raw)
    p._flush()
    text = "\n\n".join(p.out)
    return text, {"title": p.title, "owner": p.meta.get("owner") or p.meta.get("maintainer"),
                  "author": p.meta.get("author")}


HEAD_NUM = re.compile(r"^(?:section\s+)?\d+(?:\.\d+)*\.?\s+[A-Z][^.!?]{2,80}$", re.I)


def _pdf(data: bytes) -> tuple[str, dict]:
    from pypdf import PdfReader
    r = PdfReader(io.BytesIO(data))
    lines = []
    for page in r.pages:
        lines += (page.extract_text() or "").splitlines()
    out = []
    for i, l in enumerate(lines):
        s = l.strip()
        nxt = lines[i + 1].strip() if i + 1 < len(lines) else ""
        is_head = bool(HEAD_NUM.match(s)) or (
            3 < len(s) < 60 and s[0].isupper() and not s.endswith((".", ",", ":", ";")) and len(s.split()) <= 8
            and sum(w[:1].isupper() for w in s.split()) >= max(1, len(s.split()) - 2) and len(nxt) > len(s))
        if re.fullmatch(r"(page\s+)?\d+(\s+of\s+\d+)?", s, re.I):
            continue  # page numbers
        out.append(f"## {s}" if is_head else l)
    info = r.metadata or {}
    return "\n".join(out), {"title": (info.get("/Title") or None), "author": info.get("/Author")}


def _docx(data: bytes) -> tuple[str, dict]:
    import docx
    d = docx.Document(io.BytesIO(data))
    out = []
    for p in d.paragraphs:
        st = (p.style.name or "").lower()
        t = p.text.strip()
        if not t:
            continue
        if st == "title":
            out.append(f"# {t}")
        elif st.startswith("heading"):
            n = int(re.sub(r"\D", "", st) or 2)
            out.append("#" * min(max(n, 1) + 1, 3) + " " + t)
        elif "list" in st:
            out.append(f"- {t}")
        else:
            out.append(t)
    for tbl in d.tables:
        for row in tbl.rows:
            out.append(" | ".join(c.text.strip() for c in row.cells))
    cp = d.core_properties
    return "\n\n".join(out), {"title": cp.title or None, "author": cp.author or None}


def _csv(raw: str, name: str) -> tuple[str, dict]:
    rows = list(csv.reader(io.StringIO(raw)))
    if not rows:
        return "", {}
    title = name.replace("-", " ").replace("_", " ").title()
    md = [f"# {title}", "", f"## {title}", "", "| " + " | ".join(rows[0]) + " |", "|" + "---|" * len(rows[0])]
    md += ["| " + " | ".join(r) + " |" for r in rows[1:]]
    return "\n".join(md), {"title": title}


def convert(path: str, data: bytes) -> tuple[str, dict]:
    """-> (markdown, meta). meta: title, owner, owner_source ('explicit' | 'inferred' | None)."""
    ext = Path(path).suffix.lower()
    if ext == ".pdf":
        text, meta = _pdf(data)
    elif ext == ".docx":
        text, meta = _docx(data)
    else:
        raw = data.decode("utf-8", errors="replace")
        if ext in (".html", ".htm"):
            text, meta = _html(raw)
        elif ext == ".txt":
            text, meta = _txt(raw)
        elif ext == ".csv":
            text, meta = _csv(raw, Path(path).stem)
        else:
            text, meta = _md(raw)
    owner = meta.get("owner") or _owner_from_text(text)
    if owner:
        meta.update(owner=str(owner), owner_source="explicit")
    elif meta.get("author"):
        meta.update(owner=str(meta["author"]), owner_source="inferred")  # file author != accountable owner
    title = meta.get("title")
    if title and not text.lstrip().startswith(("---", "# ")):
        text = f"---\ntitle: {json_safe(title)}\n---\n" + text
    return text, meta


def json_safe(s: str) -> str:
    return '"' + str(s).replace('"', "'") + '"'
